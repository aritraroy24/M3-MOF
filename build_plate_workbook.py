"""Build plate_plan_96wells.xlsx: one sheet with the 96-well ratios, per-well dispense volumes and stock-solution preparation, all driven by Excel formulas."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from generate_plate import LIGAND_RATIOS, METALS, OVERAGE, REAGENTS, RESERVOIR_DEAD_mL, RESERVOIR_WELLS, compositions

OUT = Path(__file__).parent / "plate_plan_96wells.xlsx"
FIRST, LAST = 4, 99  # plate data rows

F = "Arial"
BOLD = Font(name=F, bold=True)
INPUT = Font(name=F, color="0000FF")
NORMAL = Font(name=F)
HEAD_FILL = PatternFill("solid", start_color="D9E1F2")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def head(ws, ref, text):
    c = ws[ref]
    c.value, c.font, c.fill, c.alignment, c.border = text, BOLD, HEAD_FILL, CENTER, BOX


def cell(ws, ref, value, font=NORMAL, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font, c.border = value, font, BOX
    c.alignment = Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Plate plan"
    ws["A1"] = "Exploring Heterometallic Co–Ni–Cu Benzenedicarboxylate Coordination Polymers via Automated Ternary Gradient Synthesis: 96-well plate plan"
    ws["A1"].font = Font(name=F, bold=True, size=14)

    # --- Plate table (left) ---
    ws.merge_cells("B2:E2")
    head(ws, "B2", "Ratio")
    ws.merge_cells("F2:J2")
    head(ws, "F2", "Volume to dispense per well (µL)")
    for col, text in zip("ABCDEFGHIJ", ["Well", "Co", "Ni", "Cu", "Ligand\n(Metal : Ligand = 1 : x)",
                                        "Co", "Ni", "Cu", "BDC", "DMF top-up"]):
        head(ws, f"{col}3", text)

    r = FIRST
    for li, L in enumerate(LIGAND_RATIOS):
        for i, (_, ratio, _) in enumerate(compositions()):
            cell(ws, f"A{r}", f"{'ABCDEFGH'[li * 2 + i // 12]}{i % 12 + 1}")
            for col, v in zip("BCD", ratio):
                cell(ws, f"{col}{r}", v, INPUT)
            cell(ws, f"E{r}", L, INPUT)
            # uL = metal umol per well x mole fraction / working conc (mM) x 1000
            for col, rc, stock_row in zip("FGH", "BCD", (11, 12, 13)):
                cell(ws, f"{col}{r}", f"=$M$4*{rc}{r}/SUM($B{r}:$D{r})/$Q${stock_row}*1000", fmt="0.00")
            cell(ws, f"I{r}", f"=$M$4*E{r}/$Q$14*1000", fmt="0.00")
            cell(ws, f"J{r}", f"=$M$3-SUM(F{r}:I{r})", fmt="0.00")
            r += 1
    assert r - 1 == LAST
    ws.conditional_formatting.add(f"J{FIRST}:J{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="FF0000", bold=True)))

    # --- Plate inputs (right) ---
    ws.merge_cells("L2:M2")
    head(ws, "L2", "Plate inputs")
    inputs = [
        ("Well volume (µL)", 300, "0", "200 µL UV-Vis aliquot (per application) + LRMS sample + losses"),
        ("Total metal per well (µmol)", 5.4, "0.00", "60 µL of 90 mM metal; sized so the 1:4 well (60 µL metal + 240 µL BDC) exactly fills 300 µL"),
        ("Working-solution overage", OVERAGE, "0.00", "Extra working solution on top of plate use + reservoir dead volume"),
    ]
    for k, (label, val, fmt, note) in enumerate(inputs, start=3):
        cell(ws, f"L{k}", label).alignment = Alignment(horizontal="left")
        c = cell(ws, f"M{k}", val, INPUT, fmt, INPUT_FILL)
        c.comment = Comment(note, "Plate plan")
    cell(ws, "L6", "Final total metal in well (mM)").alignment = Alignment(horizontal="left")
    cell(ws, "M6", "=M4/M3*1000", fmt="0.00")
    cell(ws, "L7", "Final BDC in well (mM)").alignment = Alignment(horizontal="left")
    cell(ws, "M7", f'=TEXT(M6*MIN(E{FIRST}:E{LAST}),"0")&" - "&TEXT(M6*MAX(E{FIRST}:E{LAST}),"0")')
    cell(ws, "L8", "Reservoir dead volume per trough (mL)").alignment = Alignment(horizontal="left")
    c = cell(ws, "M8", RESERVOIR_DEAD_mL, INPUT, "0.0", INPUT_FILL)
    c.comment = Comment("Axygen 12-well reservoir (8 x 71 mm troughs): ~1.1 mL lies below the 2 mm aspirate height, plus margin", "Plate plan")

    # --- Stock solutions (right) ---
    ws.merge_cells("L9:U9")
    head(ws, "L9", "Stock solutions (all in DMF)  -  weigh, dissolve, then dilute an aliquot to the working concentration")
    for col, text in zip("LMNOPQRSTU", ["Reagent", "MW\n(g/mol)", "Stock\n(mM)", "Stock volume\n(mL)", "Mass to weigh\n(mg)",
                                        "Working soln\n(mM)", "Plate needs\n(mL)", "Aliquot of stock\n(mL)", "Make up to\n(mL)",
                                        "Reservoir\ntroughs"]):
        head(ws, f"{col}10", text)
    for k, (key, plate_col) in enumerate(zip(METALS + ["BDC"], "FGHI"), start=11):
        g = REAGENTS[key]
        cell(ws, f"L{k}", f"{key}: {g['name']}").alignment = Alignment(horizontal="left")
        cell(ws, f"M{k}", g["mw"], INPUT, "0.00")
        cell(ws, f"N{k}", g["stock_mM"], INPUT, "0", INPUT_FILL)
        cell(ws, f"O{k}", g["stock_mL"], INPUT, "0", INPUT_FILL)
        cell(ws, f"P{k}", f"=N{k}*O{k}*M{k}/1000", fmt="0.0")
        cell(ws, f"Q{k}", 90, INPUT, "0", INPUT_FILL)
        cell(ws, f"R{k}", f"=SUM({plate_col}{FIRST}:{plate_col}{LAST})/1000", fmt="0.00")
        cell(ws, f"S{k}", f"=T{k}*Q{k}/N{k}", fmt="0.00")
        cell(ws, f"T{k}", f"=ROUNDUP((R{k}+$M$8*U{k})*$M$5,0)", fmt="0")
        cell(ws, f"U{k}", RESERVOIR_WELLS[key], INPUT, "0")
    cell(ws, "L15", "DMF: N,N-Dimethylformamide (well top-up)").alignment = Alignment(horizontal="left")
    for col in "MNOPQS":
        cell(ws, f"{col}15", "")
    cell(ws, "R15", f"=SUM(J{FIRST}:J{LAST})/1000", fmt="0.00")
    cell(ws, "T15", "=ROUNDUP((R15+$M$8*U15)*$M$5,0)", fmt="0")
    cell(ws, "U15", RESERVOIR_WELLS["DMF"], INPUT, "0")

    notes = [
        "Notes",
        "Blue = input you can change (yellow = key choices); black = calculated.",
        "Dilution: pipette 'Aliquot of stock' and add DMF to the 'Make up to' total volume (C1V1 = C2V2).",
        "'Make up to' covers plate use + reservoir dead volume x troughs, times the overage; split it evenly across the troughs.",
        "Masses assume 100% purity; divide by the certificate-of-analysis purity before weighing.",
        "BDC solubility in DMF is ~0.38 M at 25 °C, hence a 180 mM stock. Check Ni/Co acetate dissolve at 450 mM before weighing the full amount.",
        "DMF top-up turns red if a well would overflow the set well volume.",
    ]
    for k, text in enumerate(notes, start=17):
        ws[f"L{k}"] = text
        ws[f"L{k}"].font = BOLD if k == 17 else NORMAL

    ws.row_dimensions[3].height = 45
    ws.row_dimensions[10].height = 45
    for col, w in {"A": 7, "B": 6, "C": 6, "D": 6, "E": 13, "F": 8, "G": 8, "H": 8, "I": 8, "J": 10, "K": 3,
                   "L": 46, "M": 10, "N": 9, "O": 12, "P": 13, "Q": 12, "R": 11, "S": 15, "T": 11, "U": 10}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B4"
    wb.save(OUT)


if __name__ == "__main__":
    main()
